#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将 PPTX 逐页转成高清 PNG 截图，再合成新的纯截图版 PPTX。
用法: python3 ppt_screenshot.py <input.pptx>
"""
import fitz
import subprocess
import os
import sys
from pptx import Presentation

DPI = 300


def export_pptx_to_pdf(pptx_path, pdf_path):
    script = f'''
tell application "Microsoft PowerPoint"
    activate
    open POSIX file "{pptx_path}"
    save active presentation in POSIX file "{pdf_path}" as save as PDF
    close active presentation
end tell
'''
    subprocess.run(["osascript", "-e", script], check=True)
    print(f"[1/3] Exported PDF: {pdf_path}")


def pdf_to_images(pdf_path, output_dir, prefix):
    doc = fitz.open(pdf_path)
    image_paths = []
    for i in range(len(doc)):
        page = doc.load_page(i)
        mat = fitz.Matrix(DPI / 72, DPI / 72)
        pix = page.get_pixmap(matrix=mat)
        img_path = os.path.join(output_dir, f"{prefix}_page_{i + 1:03d}.png")
        pix.save(img_path)
        image_paths.append(img_path)
    doc.close()
    print(f"[2/3] Generated {len(image_paths)} PNG images")
    return image_paths


def create_screenshot_pptx(original_pptx, image_paths, output_pptx):
    prs_orig = Presentation(original_pptx)
    slide_width = prs_orig.slide_width
    slide_height = prs_orig.slide_height

    prs_new = Presentation()
    prs_new.slide_width = slide_width
    prs_new.slide_height = slide_height

    blank_layout = None
    for layout in prs_new.slide_layouts:
        if layout.name == "Blank":
            blank_layout = layout
            break
    if blank_layout is None:
        blank_layout = prs_new.slide_layouts[-1]

    for img_path in image_paths:
        slide = prs_new.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            img_path,
            left=0,
            top=0,
            width=slide_width,
            height=slide_height,
        )

    prs_new.save(output_pptx)
    print(f"[3/3] Saved screenshot PPTX: {output_pptx}")


def main():
    if len(sys.argv) < 2:
        print("Usage: python3 ppt_screenshot.py <input.pptx>")
        sys.exit(1)

    input_pptx = os.path.abspath(sys.argv[1])
    if not os.path.exists(input_pptx):
        print(f"File not found: {input_pptx}")
        sys.exit(1)

    output_dir = os.path.dirname(input_pptx)
    basename = os.path.splitext(os.path.basename(input_pptx))[0]
    temp_pdf = os.path.join(output_dir, f"_{basename}_temp.pdf")
    output_pptx = os.path.join(output_dir, f"{basename}_screenshots.pptx")
    temp_prefix = f"_{basename}_temp"

    try:
        export_pptx_to_pdf(input_pptx, temp_pdf)
        image_paths = pdf_to_images(temp_pdf, output_dir, temp_prefix)
        create_screenshot_pptx(input_pptx, image_paths, output_pptx)
    finally:
        if os.path.exists(temp_pdf):
            os.remove(temp_pdf)
        for f in os.listdir(output_dir):
            if f.startswith(temp_prefix) and f.endswith(".png"):
                os.remove(os.path.join(output_dir, f))
        print("[Cleanup] Temporary files removed")


if __name__ == "__main__":
    main()
